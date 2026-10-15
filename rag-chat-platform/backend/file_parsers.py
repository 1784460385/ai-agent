import json
import logging
import os
import tempfile
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from openpyxl import load_workbook
from pptx import Presentation

_log = logging.getLogger("rag.io")

from config import (
    IMAGE_SUFFIXES,
    OCR_EMBEDDED_IMAGE_LIMIT,
    OCR_ENABLED,
    PDF_OCR_MIN_CHARS,
)


class ParseError(Exception):
    """文件能保存，但无法提取可索引文本。"""


def load_bytes(filename: str, data: bytes) -> list[Document]:
    """把数据库中的文件写入临时路径后再解析，来源名用原始文件名。"""
    suffix = Path(filename).suffix.lower()
    descriptor, name = tempfile.mkstemp(suffix=suffix or ".bin")
    os.close(descriptor)
    path = Path(name)
    try:
        path.write_bytes(data)
        documents = load_file(path)
    finally:
        path.unlink(missing_ok=True)
    for document in documents:
        document.metadata["source"] = filename
    return documents


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix in IMAGE_SUFFIXES:
        text = _read_image(path)
        if not text.strip():
            raise ParseError("图片中未识别到文字")
        return [
            Document(
                page_content=text,
                metadata={"source": str(path), "content_type": "ocr_image"},
            )
        ]
    text = extract_text(path)
    if not text.strip():
        raise ParseError("未提取到文本")
    return [Document(page_content=text, metadata={"source": str(path)})]


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        return _read_image(path)
    loaders = {
        ".txt": _read_plain,
        ".md": _read_plain,
        ".markdown": _read_plain,
        ".docx": _read_docx,
        ".doc": _read_doc,
        ".xlsx": _read_excel,
        ".xls": _read_excel,
        ".csv": _read_csv,
        ".pptx": _read_pptx,
        ".html": _read_html,
        ".htm": _read_html,
        ".json": _read_json,
        ".ipynb": _read_ipynb,
    }
    loader = loaders.get(suffix)
    if loader is None:
        raise ParseError("不支持的文件格式")
    return loader(path).strip()


def _ocr_bytes(data: bytes) -> str:
    if not OCR_ENABLED:
        return ""
    from ocr import OcrError, ocr_image_bytes

    try:
        return ocr_image_bytes(data)
    except OcrError as exc:
        _log.warning("OCR 跳过: %s", exc)
        return ""


def _format_ocr_blocks(blocks: list[str], *, title: str = "图片文字") -> str:
    parts = []
    for index, text in enumerate(blocks, start=1):
        cleaned = text.strip()
        if cleaned:
            parts.append(f"## {title} {index}\n\n{cleaned}")
    return "\n\n".join(parts)


def _read_image(path: Path) -> str:
    if not OCR_ENABLED:
        raise ParseError("OCR 已禁用，无法解析图片")
    from ocr import OcrError, ocr_image_bytes

    try:
        text = ocr_image_bytes(path.read_bytes())
    except OcrError as exc:
        raise ParseError(str(exc)) from exc
    return text.strip()


def _read_plain(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gbk"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("无法识别文件编码")


def _load_pdf(path: Path) -> list[Document]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ParseError("解析失败：文件已加密")
    except ParseError:
        raise
    except Exception as exc:
        message = str(exc).lower()
        if "encrypt" in message or "password" in message:
            raise ParseError("解析失败：文件已加密") from exc
        raise ParseError(f"解析失败：{exc}") from exc
    try:
        documents = PyPDFLoader(str(path)).load()
    except Exception as exc:
        message = str(exc).lower()
        if "encrypt" in message or "password" in message:
            raise ParseError("解析失败：文件已加密") from exc
        raise ParseError(f"解析失败：{exc}") from exc

    # 扫描页 / 内嵌图：文字过少时用 OCR 补全
    if OCR_ENABLED:
        documents = _enrich_pdf_with_ocr(path, documents)

    if not any(item.page_content.strip() for item in documents):
        raise ParseError("未提取到文本")
    return documents


def _enrich_pdf_with_ocr(path: Path, documents: list[Document]) -> list[Document]:
    try:
        import fitz
    except ImportError:
        return documents

    try:
        pdf = fitz.open(str(path))
    except Exception as exc:  # noqa: BLE001
        _log.warning("PDF OCR 打开失败: %s", exc)
        return documents

    enriched: list[Document] = []
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            existing = ""
            if index < len(documents):
                existing = (documents[index].page_content or "").strip()
            page_text = existing
            if len(existing) < PDF_OCR_MIN_CHARS:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                ocr_text = _ocr_bytes(pix.tobytes("png"))
                if ocr_text:
                    page_text = ocr_text if not existing else f"{existing}\n\n{ocr_text}"
            else:
                # 正文够用时，仅补充较大内嵌图中的文字
                image_texts: list[str] = []
                for img_index, image in enumerate(page.get_images(full=True)):
                    if img_index >= OCR_EMBEDDED_IMAGE_LIMIT:
                        break
                    xref = image[0]
                    try:
                        payload = pdf.extract_image(xref)
                        blob = payload.get("image") or b""
                        if len(blob) < 8_000:
                            continue
                        text = _ocr_bytes(blob)
                        if text:
                            image_texts.append(text)
                    except Exception:
                        continue
                if image_texts:
                    page_text = f"{existing}\n\n{_format_ocr_blocks(image_texts)}"
            meta = dict(documents[index].metadata) if index < len(documents) else {"page": index}
            meta.setdefault("page", index)
            enriched.append(Document(page_content=page_text, metadata=meta))
    finally:
        pdf.close()
    return enriched or documents


def _docx_element_text(element) -> str:
    """提取 Word 最终可见文本：包含修订插入(w:ins)，排除已删除(w:del)。

    python-docx 的 paragraph.text / cell.text 不会遍历 w:ins 内的 run，
    带修订痕迹的公文常把整段职责正文放在插入标记里，导致知识库只剩条目标题。
    """
    from docx.oxml.ns import qn

    parts: list[str] = []

    def walk(node) -> None:
        for child in node.iterchildren():
            tag = child.tag
            if tag == qn("w:del"):
                continue
            if tag in (qn("w:drawing"), qn("w:pict"), qn("w:object"), qn("w:commentReference")):
                continue
            if tag == qn("w:t"):
                parts.append(child.text or "")
            elif tag == qn("w:tab"):
                parts.append("\t")
            elif tag in (qn("w:br"), qn("w:cr")):
                parts.append("\n")
            else:
                walk(child)

    walk(element)
    return "".join(parts)


def _docx_list_has_label(text: str) -> bool:
    """判断列表项是否以「角色/部门：」这类标签开头。"""
    body = text[2:].strip() if text.startswith("- ") else text.strip()
    for sep in ("：", ":"):
        if sep not in body:
            continue
        label, _rest = body.split(sep, 1)
        label = label.strip()
        if 0 < len(label) <= 40 and "\n" not in label:
            return True
    return False


def _format_docx_paragraph(paragraph) -> str | None:
    text = _docx_element_text(paragraph._element).strip()
    if not text:
        return None
    style = paragraph.style.name if paragraph.style is not None else ""
    if style.startswith("Heading"):
        level_text = style.replace("Heading", "").strip() or "1"
        try:
            level = min(int(level_text), 6)
        except ValueError:
            level = 1
        return f"{'#' * level} {text}"
    # 编号列表或带 numPr 的段落都按列表输出
    from docx.oxml.ns import qn

    p_pr = paragraph._element.find(qn("w:pPr"))
    is_list = "List" in style or (
        p_pr is not None and p_pr.find(qn("w:numPr")) is not None
    )
    if is_list:
        return f"- {text}"
    return text


def _read_docx(path: Path) -> str:
    """按正文出现顺序提取段落和表格，避免表格被整段挪到文末导致上下文断裂。"""
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = DocxDocument(str(path))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            rendered = _format_docx_paragraph(Paragraph(child, document))
            if not rendered:
                continue
            # 修订后职责正文常拆成「标题段 + 续写段」两个列表项，合并以免检索只命中空标题
            if (
                rendered.startswith("- ")
                and parts
                and parts[-1].startswith("- ")
                and _docx_list_has_label(parts[-1])
                and not _docx_list_has_label(rendered)
            ):
                cont = rendered[2:].strip()
                prev = parts[-1]
                joiner = "" if prev.endswith(("：", ":", "。", "；", ";")) else " "
                parts[-1] = f"{prev}{joiner}{cont}"
            else:
                parts.append(rendered)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            rows = []
            for row in table.rows:
                rows.append(
                    [
                        _docx_element_text(cell._tc).strip().replace("\n", " ")
                        for cell in row.cells
                    ]
                )
            markdown = _rows_to_markdown(rows)
            if markdown:
                parts.append(markdown)
    image_texts = _ocr_office_images(document.part)
    if image_texts:
        parts.append(_format_ocr_blocks(image_texts, title="文档内嵌图片"))
    return "\n\n".join(parts)


def _ocr_office_images(part) -> list[str]:
    """从 Word/PPT 包内关系中提取图片并 OCR。"""
    if not OCR_ENABLED:
        return []
    texts: list[str] = []
    try:
        rels = part.rels.values()
    except Exception:
        return []
    for rel in rels:
        if len(texts) >= OCR_EMBEDDED_IMAGE_LIMIT:
            break
        try:
            reltype = str(getattr(rel, "reltype", "") or "")
            if "image" not in reltype:
                continue
            blob = rel.target_part.blob
            if not blob or len(blob) < 2_000:
                continue
            text = _ocr_bytes(blob)
            if text:
                texts.append(text)
        except Exception:
            continue
    return texts


def _read_doc(path: Path) -> str:
    data = path.read_bytes()
    chunks: list[str] = []
    current: list[str] = []
    index = 0
    while index < len(data) - 1:
        code = data[index] + (data[index + 1] << 8)
        if 32 <= code < 127 or code in (9, 10, 13) or 0x4E00 <= code <= 0x9FFF:
            current.append(chr(code))
            index += 2
            continue
        if len(current) >= 8:
            chunks.append("".join(current).strip())
        current = []
        index += 2
    if len(current) >= 8:
        chunks.append("".join(current).strip())
    text = "\n".join(item for item in chunks if item)
    if len(text) < 8:
        raise ParseError("旧版 .doc 未能提取正文，请另存为 .docx")
    return text


def _read_excel(path: Path) -> str:
    if path.suffix.lower() == ".xls":
        return _read_xls(path)
    workbook = load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
                if any(cell is not None and str(cell).strip() for cell in row)
            ]
            rendered = _rows_to_markdown(rows)
            if rendered:
                parts.append(f"## {sheet.title}\n\n{rendered}")
    finally:
        workbook.close()
    return "\n\n".join(parts)


def _read_xls(path: Path) -> str:
    import pandas as pd

    sheets = pd.read_excel(path, sheet_name=None)
    parts: list[str] = []
    for name, frame in sheets.items():
        rows = [list(frame.columns.astype(str))]
        rows.extend(frame.fillna("").astype(str).values.tolist())
        rendered = _rows_to_markdown(rows)
        if rendered:
            parts.append(f"## {name}\n\n{rendered}")
    return "\n\n".join(parts)


def _read_csv(path: Path) -> str:
    import pandas as pd

    last_error: Exception | None = None
    for encoding in ("utf-8", "utf-8-sig", "gbk"):
        try:
            frame = pd.read_csv(path, encoding=encoding)
            break
        except UnicodeDecodeError as exc:
            last_error = exc
    else:
        raise ParseError("无法识别 CSV 编码") from last_error
    rows = [list(frame.columns.astype(str))]
    rows.extend(frame.fillna("").astype(str).values.tolist())
    return _rows_to_markdown(rows)


def _read_pptx(path: Path) -> str:
    presentation = Presentation(str(path))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines = [f"## 幻灯片 {index}"]
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text.strip()
            if text:
                lines.append(text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                if not notes.startswith("备注"):
                    notes = f"备注：{notes}"
                lines.append(notes)
        if len(lines) > 1:
            parts.append("\n".join(lines))
    image_texts = _ocr_office_images(presentation.part)
    if image_texts:
        parts.append(_format_ocr_blocks(image_texts, title="幻灯片内嵌图片"))
    return "\n\n".join(parts)


def _read_html(path: Path) -> str:
    """抽取正文，去掉导航、页脚、广告等噪声。"""
    soup = BeautifulSoup(_read_plain(path), "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "iframe", "form"]):
        tag.decompose()
    for tag in soup.find_all(["nav", "footer", "header", "aside"]):
        tag.decompose()
    noise_keys = ("nav", "menu", "sidebar", "footer", "header", "advert", "ads", "cookie", "breadcrumb")
    for tag in list(soup.find_all(True)):
        attrs = " ".join(
            str(tag.get(name, ""))
            for name in ("id", "class", "role")
            if tag.get(name)
        ).lower()
        if any(key in attrs for key in noise_keys):
            tag.decompose()
    main = soup.find("main") or soup.find("article") or soup.body or soup
    lines = [line.strip() for line in main.get_text("\n").splitlines() if line.strip()]
    return "\n".join(lines)


def _read_json(path: Path) -> str:
    payload = json.loads(_read_plain(path))
    return "\n".join(_flatten_json(payload))


def _flatten_json(value, prefix: str = "") -> list[str]:
    if isinstance(value, dict):
        lines: list[str] = []
        for key, item in value.items():
            child = f"{prefix}.{key}" if prefix else str(key)
            lines.extend(_flatten_json(item, child))
        return lines
    if isinstance(value, list):
        lines = []
        for index, item in enumerate(value):
            lines.extend(_flatten_json(item, f"{prefix}[{index}]"))
        return lines
    return [f"{prefix}: {value}"]


def _read_ipynb(path: Path) -> str:
    notebook = json.loads(_read_plain(path))
    parts: list[str] = []
    for index, cell in enumerate(notebook.get("cells", []), start=1):
        source = cell.get("source", "")
        if isinstance(source, list):
            source = "".join(source)
        source = source.strip()
        if not source:
            continue
        kind = cell.get("cell_type", "cell")
        parts.append(f"## 单元格 {index}（{kind}）\n{source}")
    return "\n\n".join(parts)


def _rows_to_markdown(rows: list[list[str]]) -> str:
    cleaned = [[str(cell).replace("|", "\\|").strip() for cell in row] for row in rows if any(str(cell).strip() for cell in row)]
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    normalized = [row + [""] * (width - len(row)) for row in cleaned]
    header = normalized[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in normalized[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)
